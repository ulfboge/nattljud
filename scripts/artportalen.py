"""Artportalen-underlag enligt Artportalens Excelmall (version 4.17, gäller fr.o.m. 3 mars 2022).

Skriver exports/artportalen_<datum>_<tid>.xlsx med:
  Så här gör du        – kort instruktion
  Fladdermöss         – exakt samma kolumner som mallens blad "Fladdermöss"
  Ryggradslösa djur   – mallens blad "Ryggradslösa djur" (vårtbitare m.fl.)
  Däggdjur (exkl.fladdermöss)
  Fåglar              – mallens blad "Fåglar" (BirdNET-bestämda fåglar)
  Granska             – alla taxa per natt och lokal, även de som inte kom med

En rad per art, natt och lokal där minst en registrering når minsta sannolikhet
(standard 0,8). Rad 1 på varje blad är kolumnrubrikerna – markera rubrikraden och
fynden, kopiera och klistra in på artportalen.se/ImportSighting.

Koordinaterna är SWEREF 99 TM, så profilens koordinatsystem i Artportalen måste
vara SWEREF99 TM. Metod, noggrannhet m.m. per lokal ställs in i data/sites.json.
"""
import os
from datetime import datetime
from statistics import median

from geo import wgs84_to_sweref99tm

# Kolumnrubriker exakt som i mallen (version 4.17). "Med-observatör" upprepas 10 gånger.
_TAIL = ["Artbestämd av", "Artbestämd av (fritext)", "Bestämningsår", "Beskrivning artbestämning",
         "Bekräftad av", "Bekräftad av (fritext)", "Bekräftelseår", "Länk till BOLD/GenBank"] + \
        ["Med-observatör"] * 10 + ["Externid", "Ej funnen"]
_PLACE_TIME = ["Lokalnamn", "Ost", "Nord", "Noggrannhet", "Diffusion", "Djup min", "Djup max",
               "Höjd min", "Höjd max", "Startdatum", "Starttid", "Slutdatum", "Sluttid",
               "Publik kommentar", "Intressant kommentar", "Privat kommentar"]
SHEETS = {
    "Fladdermöss": ["Artnamn", "Antal", "Enhet", "Ålder-Stadium", "Kön", "Aktivitet", "Metod"] + _PLACE_TIME +
                   ["Ej återfunnen", "Dölj fyndet t.o.m.", "Andrahand", "Osäker artbestämning", "Ospontan",
                    "Biotop", "Biotop-beskrivning"] + _TAIL,
    "Ryggradslösa djur": ["Artnamn", "Antal", "Enhet", "Antal substrat", "Ålder-Stadium", "Kön", "Aktivitet",
                          "Metod"] + _PLACE_TIME +
                         ["Ej återfunnen", "Dölj fyndet t.o.m.", "Andrahand", "Osäker artbestämning", "Ospontan",
                          "Biotop", "Biotop-beskrivning", "Art som substrat", "Art som substrat beskrivning",
                          "Substrat", "Substrat-beskrivning", "Offentlig samling", "Privat samling",
                          "Samlings-nummer", "Bestämningsmetod"] + _TAIL,
    "Däggdjur (exkl.fladdermöss)": ["Artnamn", "Antal", "Ålder-Stadium", "Kön", "Aktivitet", "Metod"] + _PLACE_TIME +
                                   ["Ej återfunnen", "Dölj fyndet t.o.m.", "Andrahand", "Osäker artbestämning",
                                    "Ospontan", "Biotop", "Biotop-beskrivning"] + _TAIL,
    "Fåglar": ["Artnamn", "Antal", "Ålder-Stadium", "Kön", "Aktivitet", "Metod", "Lokalnamn", "Huvudlokal",
               "Ost", "Nord", "Noggrannhet", "Diffusion", "Startdatum", "Starttid", "Slutdatum", "Sluttid",
               "Publik kommentar", "Intressant kommentar", "Privat kommentar", "Ej återfunnen", "Andrahand",
               "Osäker artbestämning", "Ospontan", "Biotop", "Biotop-beskrivning", "Artbestämd av",
               "Artbestämd av (fritext)", "Bestämningsår", "Beskrivning artbestämning", "Bekräftad av",
               "Bekräftad av (fritext)", "Bekräftelseår", "Länk till BOLD/GenBank", "Dölj fyndet t.o.m."] +
              ["Med-observatör"] * 10 + ["Externid", "Ej funnen"],
}
# BTO-grupp -> (blad, standardvärden). Värdena finns i mallens listor för respektive blad.
GROUP_SHEET = {
    "bat": ("Fladdermöss", {"Enhet": "Registreringar", "Aktivitet": "Aktiv"}),
    "bush-cricket": ("Ryggradslösa djur", {"Aktivitet": "Spel", "Metod": "Ultraljudsdetektor"}),
    # Mallen saknar ultraljudsdetektor som metod för övriga däggdjur – lämnas tom.
    "terrestrial mammal": ("Däggdjur (exkl.fladdermöss)", {"Aktivitet": "Lockläte, övriga läten"}),
    # Fåglar: antal individer går inte att avgöra från ljudet – Antal lämnas tomt.
    "bird": ("Fåglar", {"Aktivitet": "Lockläte, övriga läten", "Metod": "Passiv ljudinspelning"}),
}
# Nattflyttare: Aktivitet blir Sträckande när minst hälften av klippen är från natten (kl. 19–06).
NIGHT_MIGRANTS = {"Turdus iliacus", "Turdus philomelos", "Turdus pilaris", "Fringilla montifringilla",
                  "Anas crecca"}
NOGGRANNHET = [1, 5, 10, 25, 50, 75, 100, 125, 150, 200, 250, 300, 400, 500, 750, 1000, 1500, 2000, 2500, 3000, 5000]
CALL_SV = {"echolocation": "ekolod", "social": "sociala läten"}


def _n(k):
    return f"{k} registrering" if k == 1 else f"{k} registreringar"


def _p(v):
    return f"{v:.2f}".replace(".", ",")


def _span(times):
    a, b = times[0][11:16], times[-1][11:16]
    return f"kl. {a}" if a == b else f"kl. {a}–{b}"


def _accuracy(m):
    return f"{min(NOGGRANNHET, key=lambda v: abs(v - (m or 50)))} m"


def export(rows, species, sites, out_dir, min_prob=0.8, classifier="", recorders=(), prev_ids=(), scope=""):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("  ! openpyxl saknas – kör: pip install openpyxl")
        return None

    species, rows = _merge_pairs(species, rows)
    groups = {}
    for d in rows:  # d = [spIdx, siteIdx, time, prob, callType, night, file]
        groups.setdefault((d[1], d[5], d[0]), []).append(d)

    per_sheet = {k: [] for k in SHEETS}
    granska = []
    ids = []
    for (si, night, spi), dets in sorted(groups.items(), key=lambda kv: (kv[0][1], kv[0][0], kv[0][2])):
        sp, site = species[spi], sites[si]
        probs = [d[3] for d in dets]
        good = [d for d in dets if d[3] >= min_prob]
        med = median(probs)
        is_species = sp.get("rank") == "species"
        target = GROUP_SHEET.get(sp.get("group"))
        stat = {"Natt": night, "Lokal": site["name"], "Artnamn": sp.get("sv") or sp["sci"],
                "Vetenskapligt namn": sp["sci"], "Grupp": sp.get("groupSv", ""), "Registreringar": len(dets),
                f"≥ {min_prob:g}": len(good), "Median p": round(med, 2), "Max p": round(max(probs), 2)}
        if not is_species:
            stat["Med i underlaget"] = "Nej – ej bestämd till art"
        elif not target:
            stat["Med i underlaget"] = "Nej – artgruppen hanteras inte"
        elif not good:
            stat["Med i underlaget"] = f"Nej – ingen registrering ≥ {min_prob:g}"
        else:
            stat["Med i underlaget"] = f"Ja – bladet {target[0]}"
        stat["Att tänka på"] = ("Enstaka registrering – kontrollera i spektrogram" if is_species and len(good) == 1 else
                                "Median under gränsen – överväg Osäker artbestämning" if is_species and good and med < min_prob else "")
        if sp.get("group") == "terrestrial mammal" and good:
            stat["Att tänka på"] = (stat["Att tänka på"] + "; " if stat["Att tänka på"] else "") + \
                "BTO varnar: alla näbbmusarter ingår inte i klassificeraren"
        base_id = f"nattljud:{site['id']}:{night}:{sp['sci'].replace(' ', '_')}"
        if sp.get("pairNote"):
            stat["Att tänka på"] = (stat["Att tänka på"] + "; " if stat["Att tänka på"] else "") + sp["pairNote"]
        if base_id in prev_ids:
            stat["Att tänka på"] = (stat["Att tänka på"] + "; " if stat["Att tänka på"] else "") + \
                "Arten har redan rapporterats denna natt i en tidigare export – detta är bara nya inspelningar"
        granska.append(stat)
        if not (is_species and target and good):
            continue

        sheet, defaults = target
        times = sorted(d[2] for d in good)
        e, n = wgs84_to_sweref99tm(site["lat"], site["lon"])
        calls = sorted({d[4] for d in good if d[4]})
        used = [recorders[i] for i in sorted({d[7] for d in good if len(d) > 7 and d[7] >= 0})]
        rec_txt = f" Inspelare: {', '.join(r['name'] for r in used)}." if used else ""
        rec = dict(defaults)
        bird = sheet == "Fåglar"
        if bird and sp["sci"] in NIGHT_MIGRANTS and \
                sum(not "06:00" <= t[11:16] < "19:00" for t in times) * 2 >= len(times):
            rec["Aktivitet"] = "Sträckande"
        if sheet == "Fladdermöss":
            rec["Antal"] = len(good)
            rec["Metod"] = (site.get("metod_fladdermoss") or
                            next((r.get("artportalen_metod_fladdermoss") for r in used if r.get("artportalen_metod_fladdermoss")), None)
                            or "Autobox")
        rec.update({
            "Artnamn": sp.get("apName") or sp.get("sv") or sp["sci"],
            "Lokalnamn": site["name"],
            "Ost": e, "Nord": n,
            "Noggrannhet": _accuracy(site.get("accuracy_m", 50)),
            "Startdatum": times[0][:10], "Starttid": times[0][11:16],
            "Slutdatum": times[-1][:10], "Sluttid": times[-1][11:16],
            "Publik kommentar": (f"Ljudinspelning, hörd i {len(good)} klipp à 40 s {_span(times)}." if bird else
                                 f"Ultraljud, {_n(len(good))} {_span(times)}"
                                 f"{' (' + ', '.join(CALL_SV.get(c, c) for c in calls) + ')' if calls else ''}."),
            "Privat kommentar": f"{_n(len(dets))} totalt denna natt, median sannolikhet {_p(med)}.{rec_txt}",
            "Beskrivning artbestämning": (f"Automatisk artbestämning med BirdNET 2.4 (passiv inspelare i fågelläge), "
                                          f"högsta sannolikhet {_p(max(probs))}. Ej manuellt verifierad." if bird else
                                          f"Automatisk artbestämning i BTO Acoustic Pipeline ({classifier}), "
                                          f"högsta sannolikhet {_p(max(probs))}. Ej manuellt verifierad."),
        })
        # Artportalen vägrar Externid när inget projekt är valt vid import (SiteConversionError_ExternalId_
        # NoProjectSelected) – kolumnen lämnas tom och id:t läggs i privat kommentar i stället.
        ext_id = _unique_id(base_id, prev_ids)
        rec["Privat kommentar"] += f" Id: {ext_id}"
        ids.append(ext_id)
        per_sheet[sheet].append(rec)

    wb = Workbook()
    head_font, head_fill = Font(bold=True), PatternFill("solid", fgColor="E1E0D9")
    ws = wb.active
    ws.title = "Så här gör du"
    for line in [
        "Import till Artportalen",
        "",
        "1. Kontrollera att koordinatsystemet under Min profil i Artportalen är SWEREF99 TM och artnamnsspråk Svenska.",
        "2. Gå igenom bladet Granska – särskilt enstaka registreringar och kolumnen 'Att tänka på'.",
        "3. Öppna ett artgruppsblad, markera rubrikraden (rad 1) och fynden, kopiera (Ctrl+C).",
        "4. Klistra in på artportalen.se/ImportSighting och klicka Importera. Ett blad i taget, max 2000 rader.",
        "5. Fynden hamnar under fliken Granska & publicera i Artportalen – kontrollera och publicera där.",
        "",
        "Kolumnerna följer Artportalens Excelmall version 4.17. Externid lämnas tom (kräver projekt); fyndets id står i Privat kommentar.",
        f"Minsta sannolikhet för att komma med: {min_prob:g}. Skapad {datetime.now():%Y-%m-%d %H:%M}.",
    ] + ([f"Omfattar {scope}. Exporterade filer loggas i data/artportalen_exporterat.json."] if scope else []):
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 120

    for sheet, cols in SHEETS.items():
        data = per_sheet[sheet]
        if not data:
            continue
        ws = wb.create_sheet(sheet)
        ws.append(cols)
        for r in data:
            ws.append([r.get(c) if c != "Med-observatör" else None for c in cols])
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
        for i, c in enumerate(cols, 1):
            vals = [len(str(r.get(c) or "")) for r in data] if c != "Med-observatör" else [0]
            ws.column_dimensions[get_column_letter(i)].width = min(50, max(len(c), *vals) + 2)
        # text i datum/tid så att Excel inte gör om dem vid kopiering
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if cols[c.column - 1] in ("Startdatum", "Starttid", "Slutdatum", "Sluttid"):
                    c.number_format = "@"
        ws.freeze_panes = "B2"

    ws = wb.create_sheet("Granska")
    cols = list(granska[0].keys()) if granska else ["Inga rader"]
    ws.append(cols)
    for r in granska:
        ws.append([r.get(c) for c in cols])
    for c in ws[1]:
        c.font, c.fill = head_font, head_fill
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(len(c), *(len(str(r.get(c) or "")) for r in granska)) + 2)
    ws.freeze_panes = "A2"

    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, f"artportalen_{datetime.now():%Y-%m-%d_%H%M}.xlsx")
    wb.save(path)
    _write_tsv(path[:-5], per_sheet, granska)
    return path, sum(len(v) for v in per_sheet.values()), len(granska), ids


# Arter som inte går att skilja på ljudet rapporteras som artpar (Artportalens valideringsregel för
# fladdermöss: "Om individen inte går att artbestämma, rapportera den som mys/bra"). Namn enligt Dyntaxa.
PAIRS = {
    "Myotis mystacinus": "pair_mysbra", "Myotis brandtii": "pair_mysbra",
}
PAIR_TAXA = {
    "pair_mysbra": {"sci": "Myotis mystacinus/brandtii", "sv": "Mustasch-/tajgafladdermus",
                    "apName": "mustaschfladdermus/tajgafladdermus", "dyntaxaId": 232474,
                    "en": "Whiskered/Brandt's Bat",
                    "note": "Mustaschfladdermus och tajgafladdermus går inte att skilja åt på ekolodsljud. "
                            "BTO:s klassificerare anger ändå den ena eller andra arten – här visas de "
                            "som artpar, så som de ska rapporteras till Artportalen.",
                    "pairNote": "Mustasch- och tajgafladdermus skiljs inte på ljudet – rapporteras som artpar"},
}


def _merge_pairs(species, rows):
    """Byt mustasch- och tajgafladdermus mot artparet; registreringar samma natt slås ihop till ett fynd."""
    species = list(species)
    new_idx, remap = {}, {}
    for i, sp in enumerate(species):
        key = PAIRS.get(sp.get("sci"))
        if not key:
            continue
        if key not in new_idx:
            species.append({**sp, **PAIR_TAXA[key], "rank": "species"})
            new_idx[key] = len(species) - 1
        remap[i] = new_idx[key]
    if not remap:
        return species, rows
    return species, [[remap.get(d[0], d[0])] + list(d[1:]) for d in rows]


def _unique_id(base, prev_ids):
    """Samma art, natt och lokal i en senare export får ett löpnummer så att Externid förblir unikt."""
    if base not in prev_ids:
        return base
    k = 2
    while f"{base}:{k}" in prev_ids:
        k += 1
    return f"{base}:{k}"


def _txt(v):
    return "" if v is None else str(v).replace("\t", " ").replace("\r", " ").replace("\n", " ")


def _write_tsv(folder, per_sheet, granska):
    """Samma innehåll som textfiler (tabbseparerade) – går att klistra in i Artportalen utan Excel:
    öppna i Anteckningar, Ctrl+A, Ctrl+C, klistra in på artportalen.se/ImportSighting."""
    os.makedirs(folder, exist_ok=True)
    for i, (sheet, cols) in enumerate(SHEETS.items(), 1):
        data = per_sheet[sheet]
        if not data:
            continue
        lines = ["\t".join(cols)] + ["\t".join(_txt(r.get(c)) if c != "Med-observatör" else "" for c in cols) for r in data]
        with open(os.path.join(folder, f"{i} {sheet} ({len(data)} fynd).txt"), "w", encoding="utf-8", newline="\r\n") as fh:
            fh.write("\n".join(lines) + "\n")
    if granska:
        cols = list(granska[0].keys())
        with open(os.path.join(folder, "0 Granska.txt"), "w", encoding="utf-8", newline="\r\n") as fh:
            fh.write("\n".join(["\t".join(cols)] + ["\t".join(_txt(r.get(c)) for c in cols) for r in granska]) + "\n")
